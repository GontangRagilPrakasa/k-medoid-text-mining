from flask import Flask, render_template, request
from app import app
import os
import pandas as pd
from openpyxl import load_workbook
import numpy as np
from sklearn.metrics.pairwise import pairwise_distances
from .custom_libraries import Engine, preprocess, kMedoids
from app.module.models import *
pd.set_option('display.max_colwidth', -1)

@app.route('/', methods=['GET', 'POST'])
def home():
	return render_template("index.html")

@app.route('/upload', methods=['GET', 'POST'])
def upload():
	if request.method == 'POST':
		global dataset
		global dataset_name
		dataload = request.files['file']
		dataload.save(os.path.join('app/data', 'data.xlsx'))
		wb = load_workbook(filename='app/data/data.xlsx')
		dataset = pd.DataFrame(wb['Sheet1'].values)
		dataset.columns = ['Judul']

		# Save dataset to database
		count_dataset = Dataset.query.count()
		dataset_name = "Dataset_{}".format(count_dataset + 1)
		try:
			data_model = Dataset(name=dataset_name)
			db.session.add(data_model)
			db.session.commit()
		except Exception as e:
			print("Failed to add data | {}".format(e))
		print(dataset)
		return render_template("uploaded_files.html", tables=[dataset.to_html(classes="table mb-0 border-0 table-responsive", justify='unset').replace('border="1"', "")])
	else:
		return render_template("upload.html")

@app.route('/preprocessing', methods=['GET', 'POST'])
def preprocessing():
	global dataset
	global pre_judul
	list_pre_judul = []
	try:
		for data in dataset["Judul"]:
			list_pre_judul.append(preprocess(data))
	except Exception as e:
		list_pre_judul = ["Data belum di upload"]
	pre_judul = pd.DataFrame(list_pre_judul)
	pre_judul.columns = ["Judul"]
	print(dataset)
	return render_template("preprocessing.html", tables=[pre_judul.to_html(classes="table mb-0 border-0 table-responsive", justify='unset').replace('border="1"', "")])

@app.route('/cosine', methods=['GET', 'POST'])
def cosine():
	global pre_judul
	global titles_score
	# Call Engine from Custom Libraries
	engine = Engine()
	# Define data uji
	list_dokumen = [str(x) for x in pre_judul['Judul']]
	list_datauji = [str(x) for x in pre_judul['Judul']]
	columnsName = []
	for i, doc in enumerate(list_dokumen):
		engine.addDocument(doc)
		columnsName.append("Document_{}".format(i + 1))
	for doc in list_datauji:
		engine.setQuery(doc)
	titles_score = engine.process_score()
	titlesScoreDf = pd.DataFrame(titles_score)
	titlesScoreDf.columns = columnsName
	return render_template('cosine.html', tables=[titlesScoreDf.to_html(classes="table mb-0 border-0 table-responsive", justify='unset').replace('border="1"', "")])

@app.route('/result', methods=['GET','POST'])
def result():
	if request.method == "POST":
		global titles_score
		global dataset
		k = int(request.form['k'])
		data = np.array(titles_score)
		datasetResult = dataset.copy()
		datasetResult['Cluster'] = 0
		# Get distance matrix
		D = pairwise_distances(data, metric='euclidean')
		# Get medoids point
		M, C = kMedoids(D, k)
		data_model = Dataset.query.filter_by(name=dataset_name).first()
		try:
			data_model.medoid = str(M)
			db.session.commit()
		except Exception as e:
			print("Failed to update medoid: {}".format(e))
		print("<Medoids:{}>".format(M))
		hasil = []
		for label in C:
			for point_idx in C[label]:
				datasetResult["Cluster"][point_idx] = label

		count_clusterdb = Cluster.query.count()
		clusterdb_name = "Clustering_{}".format(count_clusterdb + 1)

		try:
			clusterdb = Cluster(name=clusterdb_name, number_of_cluster=len(C), dataset_id=data_model.id)
			db.session.add(clusterdb)
			db.session.commit()
		except Exception as e:
			print("Failed to add data | {}".format(e))

		# Get Cluster Id from database filter by name
		clusterdb = Cluster.query.filter_by(name=clusterdb_name).first()

		for i in range(len(datasetResult["Cluster"])):
			cluster_data = ClusterData(title=datasetResult["Judul"][i], cluster=int(datasetResult["Cluster"][i]), cluster_id=clusterdb.id)
			try:
				db.session.add(cluster_data)
				db.session.commit()
			except Exception as e:
				print("Failed to add data | {}".format(e))
		print(datasetResult)
		return render_template('result.html', tables=[datasetResult.to_html(classes="table mb-0 border-0 table-responsive", justify='unset').replace('border="1"', "")])

@app.route("/prev-result")
def prevResult():
	# Get dataset from database
	dataset_db = Dataset.query.all()
	dataset_name = []
	for data in dataset_db:
		dataset_name.append(data.name)

	# Get clusters form database
	clusters = Cluster.query.all()
	cluster_name = []
	clusterDf = []
	listLenCluster = []
	obj = {}
	for cluster in clusters:
		listClusterDataDf = []
		cluster_name.append(cluster.name)
		listLenCluster.append(cluster.number_of_cluster)
		for i in range(cluster.number_of_cluster):
			cluster_data = ClusterData.query.filter_by(cluster=i, cluster_id=cluster.id).all()
			print(cluster.id)
			listClusterData = []
			for data in cluster_data:
				obj = {
					'Judul': data.title
				}
				listClusterData.append(obj)
			listClusterDataDf.append(pd.DataFrame(listClusterData, columns=obj.keys()))
		clusterDf.append(listClusterDataDf)
		print(clusterDf)

	return render_template("prev-result.html", clusterDf=clusterDf, cluster_name=cluster_name, dataset_name=dataset_name, lenDataset=len(dataset_name), listLenCluster=listLenCluster)